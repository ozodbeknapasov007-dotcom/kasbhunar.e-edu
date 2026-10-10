import json
import os
import argparse
import zipfile
import qrcode
from PIL import Image
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

def generate_all(base_domain="https://kasbhunare-edu.vercel.app"):
    with open('students.json', 'r', encoding='utf-8') as f:
        students = json.load(f)

    base_domain = base_domain.rstrip('/')
    out_dir = "qrcodes"
    os.makedirs(out_dir, exist_ok=True)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "QR Havolalar"

    headers = [
        "T/R", 
        "F.I.O", 
        "Sertifikat raqami", 
        "Qayd raqami", 
        "JShShIR (PINFL)", 
        "Tug'ilgan sana",
        "Ta'lim muassasasi",
        "Mutaxassislik",
        "O'quv soati",
        "O'qish davri",
        "DAK qarori",
        "DAK sanasi",
        "Jonli QR Havola", 
        "QR Fayl nomi"
    ]
    ws.append(headers)

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    print(f"Baza domen: {base_domain}")
    print(f"{len(students)} ta yuqori sifatli (bosmaga tayyor 300 DPI) QR kod yaratilmoqda...")

    for s in students:
        idx = s['id']
        token = s['token']
        fio = s['fio']
        cert_no = s['sertifikat_raqami']
        reg_no = s['qayd_raqami']
        pinfl = s['pinfl']
        bdate = s['tugilgan_sana']
        tm = s['tm']
        mutaxassislik = s['mutaxassislik']
        soat = s['oqish_soati']
        davr = s['oqish_davri']
        dak_q = s['dak_qarori']
        dak_s = s['dak_sanasi']
        
        qr_url = f"{base_domain}/prof/pt/{token}"
        safe_name = f"{idx:02d}_{s['familiya']}_{s['ismi']}_{reg_no}.png".replace("'", "").replace(" ", "_")
        qr_path = os.path.join(out_dir, safe_name)

        # Ultra sharp 300 DPI QR code
        qr = qrcode.QRCode(
            version=None,
            error_correction=qrcode.constants.ERROR_CORRECT_M,
            box_size=16,
            border=3
        )
        qr.add_data(qr_url)
        qr.make(fit=True)

        img = qr.make_image(fill_color="black", back_color="white")
        img.save(qr_path, dpi=(300, 300))

        ws.append([
            idx,
            fio,
            cert_no,
            reg_no,
            pinfl,
            bdate,
            tm,
            mutaxassislik,
            soat,
            davr,
            dak_q,
            dak_s,
            qr_url,
            safe_name
        ])

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    excel_report = "TALABALAR_QR_HAVOLALARI.xlsx"
    wb.save(excel_report)

    # Archive to ZIP
    zip_name = f"QR_KODLAR_{len(students)}_TALABA.zip"
    with zipfile.ZipFile(zip_name, 'w', zipfile.ZIP_DEFLATED) as zipf:
        for file in os.listdir(out_dir):
            if file.endswith('.png'):
                zipf.write(os.path.join(out_dir, file), file)

    print(f"[OK] Barcha {len(students)} ta QR kod '{out_dir}/' papkasiga saqlandi!")
    print(f"[OK] Chiroyli hisobot jadvali: '{excel_report}'")
    print(f"[OK] Barcha QR kodlar arxivda: '{zip_name}'")

if __name__ == '__main__':
    parser = argparse.ArgumentParser(description="Farmatsiya talabalari uchun xavfsiz QR kodlar generatori")
    parser.add_argument("--domain", type=str, default="https://kasbhunare-edu.vercel.app", help="Saytingizning domeni")
    args = parser.parse_args()

    generate_all(args.domain)
